import io

import openpyxl

from gfs_analysis import classify, db, export, ingest
from gfs_analysis import quality as q
from gfs_analysis.sample import make_sample_frame, write_survey_excel


def test_build_workbook(conn, tmp_path):
    ingest.import_survey_file(conn, write_survey_excel(tmp_path / "s.xlsx", make_sample_frame(120, seed=2)))
    classify.auto_classify(conn)
    all_rows = db.load_responses(conn)
    filtered = q.apply_dedup(all_rows, q.DEDUP_NEAR, 60)
    data = export.build_workbook(filtered, all_rows, {"期間": "全期間"}, db.load_issues(conn), 60)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for sheet in ["概要", "NPS", "満足度", "変化評価", "変化評価×NPS", "変化理由分類_集計",
                  "変化理由分類_明細", "品質", "品質_不正データ", "明細"]:
        assert sheet in wb.sheetnames
    assert wb["明細"].max_row == len(filtered) + 1
    texts = [c.value for r in wb["概要"].iter_rows() for c in r if isinstance(c.value, str)]
    assert any("因果関係" in t for t in texts)
    # 会員ID は文字列のまま出力される
    header = [c.value for c in wb["明細"][1]]
    col = header.index("会員ID") + 1
    assert isinstance(wb["明細"].cell(row=2, column=col).value, str)


def test_build_workbook_with_actions(conn, tmp_path):
    from gfs_analysis import actions
    from gfs_analysis.sample import action_csv_bytes, link_csv_bytes, make_action_frame
    survey = make_sample_frame(120, seed=7)
    ingest.import_survey_file(conn, write_survey_excel(tmp_path / "s.xlsx", survey))
    actions.import_id_links(conn, link_csv_bytes(survey["会員ID"].unique()), "l.csv")
    actions.import_actions(conn, action_csv_bytes(make_action_frame(survey, seed=7)), "a.csv")
    all_rows = db.load_responses(conn)
    ad = actions.attach_actions(all_rows, actions.load_action_snapshots(conn))
    data = export.build_workbook(all_rows, all_rows, {}, db.load_issues(conn), 60, ad, actions.TIMING_BEFORE)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "アクション×NPS" in wb.sheetnames

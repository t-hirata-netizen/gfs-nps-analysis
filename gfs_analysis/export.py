"""集計結果・分類結果の Excel 出力"""
from __future__ import annotations

import io
from typing import Dict, List, Optional, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from . import metrics as m
from . import quality as q

Tables = List[Tuple[str, pd.DataFrame]]


def _write_tables(writer: pd.ExcelWriter, sheet: str, tables: Tables, index: bool = False):
    row = 0
    for title, df in tables:
        pd.DataFrame([[title]]).to_excel(writer, sheet_name=sheet, startrow=row,
                                         index=False, header=False)
        writer.sheets[sheet].cell(row=row + 1, column=1).font = Font(bold=True)
        df.to_excel(writer, sheet_name=sheet, startrow=row + 1, index=index)
        row += len(df) + 4


def _autofit(ws, max_width: int = 60):
    for col in ws.columns:
        width = 8
        for cell in col:
            if cell.value is None:
                continue
            text = str(cell.value)
            w = sum(2 if ord(ch) > 0x2000 else 1 for ch in text[:200])
            width = max(width, min(w + 2, max_width))
            if len(text) > 40:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


def _summary_table(d: Dict) -> pd.DataFrame:
    return pd.DataFrame(
        [{"項目": k, "値": (round(v, 1) if isinstance(v, float) else v)} for k, v in d.items()]
    )


def detail_frame(df: pd.DataFrame) -> pd.DataFrame:
    cols = ["response_id", "member_id", "answered_at", "gfs_satisfaction", "support_label",
            "nps", "nps_category", "change_label", "admission_good", "re_enroll"] + m.TEXT_COLS
    return df[cols].rename(columns=m.COLUMN_LABELS)


def classification_frame(df: pd.DataFrame) -> pd.DataFrame:
    d = df[df["reason_change"].notna()]
    out = pd.DataFrame({
        "回答ID": d["response_id"],
        "会員ID": d["member_id"],
        "回答日時": d["answered_at"],
        "変化評価": d["change_label"],
        "NPS区分": d["nps_category"],
        "変化理由": d["reason_change"],
    })
    for c in m.CATEGORY_COLS:
        out[m.CATEGORY_LABELS[c]] = d[c].map({1: "○", 0: ""})
    out["分類状態"] = d["class_source"].map(
        {"auto": "自動（未確認）", "ai": "AI確認（人の確認待ち）", "manual": "確認済み"}).fillna("未分類")
    out["更新日時"] = d["class_updated_at"]
    return out.reset_index(drop=True)


def build_workbook(filtered: pd.DataFrame, all_rows: pd.DataFrame, conditions: Dict,
                   issues: pd.DataFrame, threshold_sec: float,
                   action_df: Optional[pd.DataFrame] = None, action_timing: str = "") -> bytes:
    """filtered: 絞り込み・重複処理後のデータ / all_rows: 品質チェック用の全データ"""
    nps = m.nps_summary(filtered)
    notes = pd.DataFrame({"注意事項": [
        m.CORRELATION_NOTE,
        "NPS = 推奨者率(9〜10) − 批判者率(0〜6)。0〜10 の有効回答のみを分母にしています。",
        "未回答・設問なし（設問追加前の回答）は 0 点や「変化なし」として数えていません。",
        f"有効回答が {m.SMALL_N} 件未満の集計は「少数データ」として扱い、解釈に注意してください。",
        "変化理由の分類は複数該当を認めているため、分類別件数の合計は回答数と一致しません。",
    ]})
    cond = pd.DataFrame([{"条件": k, "内容": v} for k, v in conditions.items()])

    change_by = m.group_nps(filtered, "change_label", m.CHANGE_OPTIONS)
    ct_counts, ct_pct = m.crosstab(filtered, "change_label", "nps_category",
                                   m.CHANGE_OPTIONS, m.NPS_CATEGORY_OPTIONS)
    exploded = m.explode_categories(filtered)
    cat_nps = m.group_nps(exploded, "change_category", m.category_order())

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        _write_tables(w, "概要", [("集計条件", cond), ("NPS サマリー", _summary_table(nps)),
                                 ("注意事項", notes)])
        _write_tables(w, "NPS", [
            ("NPS 0〜10 の分布", m.distribution(filtered, "nps_label", m.NPS_VALUES, [m.MISSING, m.NOT_ASKED])),
            ("NPS 区分", m.distribution(filtered, "nps_category", m.NPS_CATEGORIES, [m.MISSING, m.NOT_ASKED])),
        ])
        _write_tables(w, "満足度", [
            ("GFS満足度", m.distribution(filtered, "gfs_label", m.SCORE5_VALUES, [m.MISSING])),
            ("サポート満足度", m.distribution(filtered, "support_label", m.SCORE5_VALUES, [m.NOT_USED, m.MISSING])),
        ])
        _write_tables(w, "変化評価", [
            ("変化評価（1〜5、5=強く変化を感じる）",
             m.distribution(filtered, "change_label", m.SCORE5_VALUES, [m.MISSING, m.NOT_ASKED])),
        ])
        _write_tables(w, "変化評価×NPS", [
            ("変化評価別の NPS（相関であり因果ではありません）", change_by),
            ("件数クロス表（行: 変化評価 / 列: NPS区分）", ct_counts.reset_index()),
            ("行％クロス表", ct_pct.reset_index()),
        ])
        _write_tables(w, "変化理由分類_集計", [
            ("分類別の該当件数（複数該当あり）", m.category_summary(filtered)),
            ("分類別の NPS（相関であり因果ではありません）", cat_nps),
        ])
        classification_frame(filtered).to_excel(w, sheet_name="変化理由分類_明細", index=False)
        _write_tables(w, "品質", [
            ("複数回答のある会員（全データ）", q.duplicate_members(all_rows)),
            (f"近接重複（{threshold_sec:g}秒以内の連続回答）", q.near_duplicate_table(all_rows, threshold_sec)),
            ("欠損の状況（絞り込み後）", q.missing_report(filtered)),
        ])
        issues_out = issues.rename(columns={
            "batch_id": "取込ID", "file_name": "ファイル名", "source_row_no": "Excel行番号",
            "member_id": "会員ID", "field": "列", "raw_value": "元の値",
            "issue_type": "内容", "action": "処理"})
        issues_out.to_excel(w, sheet_name="品質_不正データ", index=False)
        if action_df is not None and action_df["has_action"].any():
            from . import actions
            _write_tables(w, "アクション×NPS", [
                ("使用したアクションリスト回答", pd.DataFrame({"条件": [action_timing]})),
                ("注意事項", pd.DataFrame({"注意事項": actions.ANALYSIS_NOTES})),
                ("紐づけの状況", actions.coverage(action_df)),
                ("相関（スピアマンの順位相関）", actions.correlation_table(action_df)),
                ("達成項目数ごとの NPS・変化評価", actions.progress_table(action_df)),
                ("タスク別：達成／未達成の比較", actions.task_table(action_df)),
            ])
        detail_frame(filtered).to_excel(w, sheet_name="明細", index=False)
        for ws in w.book.worksheets:
            _autofit(ws)
    return buf.getvalue()
